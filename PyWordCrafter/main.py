import json
from docx import Document
import unicodedata

def sanitize_text(text):
    return unicodedata.normalize("NFKD", text).encode("latin-1", "ignore").decode("latin-1")

# Load sections from JSON
with open("report_sections.json", "r", encoding="utf-8") as f:
    sections = json.load(f)

# Create a Word Document
doc = Document()

# Add Title
doc.add_heading("Technical Research Report: Backend Architecture for Mobile App", 0)

# Add sections
for section in sections:
    doc.add_heading(sanitize_text(section["title"]), level=1)
    doc.add_paragraph(sanitize_text(section["content"]))

# Save the document
output_path = "Backend_Architecture_Report.docx"
doc.save(output_path)

print(f"Report saved to {output_path}")
