# import json
# from fpdf import FPDF
# import unicodedata

# def sanitize_text(text):
#     return unicodedata.normalize("NFKD", text).encode("latin-1", "ignore").decode("latin-1")

# class PDF(FPDF):
#     def header(self):
#         self.set_font("Arial", "B", 12)
#         self.cell(0, 10, "Technical Research Report: Backend Architecture for Mobile App", 0, 1, "C")

#     def chapter_title(self, title):
#         self.set_font("Arial", "B", 11)
#         self.cell(0, 10, sanitize_text(title), 0, 1, "L")
#         self.ln(2)

#     def chapter_body(self, body):
#         self.set_font("Arial", "", 10)
#         self.multi_cell(0, 8, body)
#         self.ln()

#     def add_section(self, title, content):
#         self.chapter_title(title)
#         self.chapter_body(sanitize_text(content))

# # Load sections from JSON
# with open("report_sections.json", "r", encoding="utf-8") as f:
#     sections = json.load(f)

# # Initialize PDF
# pdf = PDF()
# pdf.add_page()
# pdf.set_auto_page_break(auto=True, margin=15)

# # Add sections to PDF
# for section in sections:
#     pdf.add_section(section["title"], section["content"])

# # Save PDF
# output_path = "Backend_Architecture_Report.pdf"
# pdf.output(output_path)

# print(f"Report saved to {output_path}")


import json
from docx import Document
import unicodedata

def sanitize_text(text):
    return unicodedata.normalize("NFKD", text).encode("latin-1", "ignore").decode("latin-1")

# Load sections from JSON
with open("report_sections.json", "r", encoding="utf-8") as f:
    sections = json.load(f)

# Create a Word Document
doc = Document()

# Add Title
doc.add_heading("Technical Research Report: Backend Architecture for Mobile App", 0)

# Add sections
for section in sections:
    doc.add_heading(sanitize_text(section["title"]), level=1)
    doc.add_paragraph(sanitize_text(section["content"]))

# Save the document
output_path = "Backend_Architecture_Report.docx"
doc.save(output_path)

print(f"Report saved to {output_path}")
